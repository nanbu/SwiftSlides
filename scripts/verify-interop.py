#!/usr/bin/env python3
"""Independent python-pptx checks. LibreOffice resaving is mandatory when --libreoffice is set."""
import argparse, shutil, subprocess, tempfile, zipfile
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--libreoffice', action='store_true')
parser.add_argument('--update-fixture', action='store_true')
args=parser.parse_args()
root=Path(__file__).resolve().parents[1]
with tempfile.TemporaryDirectory(prefix='swiftslides-interop-') as folder:
 temp=Path(folder); proposal=temp/'proposal.pptx'
 subprocess.run(['swift','run','--package-path',str(root),'swiftslides','sample',str(proposal)],check=True)
 with zipfile.ZipFile(proposal) as archive:
  masterParts = [n for n in archive.namelist() if n.startswith("ppt/notesMasters/") and n.endswith(".xml")]
  assert len(masterParts)==1
  main = etree.fromstring(archive.read("ppt/presentation.xml"))
  assert main.find("{*}notesMasterIdLst") is None
  for n in archive.namelist():
   if n.startswith("ppt/") and n.endswith(".xml"):
    xml = etree.fromstring(archive.read(n))
    assert not xml.xpath("//a:t/@xml:space",namespaces={"a":"http://schemas.openxmlformats.org/drawingml/2006/main"})
 p=Presentation(proposal)
 assert len(p.slides)==2 and p.slide_width/12700==960 and p.slide_height/12700==540
 assert p.slides[0].shapes[0].text=='収益性を高める3つの施策'
 assert p.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt==30
 assert p.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.bold
 assert p.slides[0].shapes[3].shape_type==MSO_SHAPE_TYPE.AUTO_SHAPE
 assert p.slides[0].shapes[3].auto_shape_type.name=='ROUNDED_RECTANGLE'
 assert p.slides[1].shapes[4].shape_type==MSO_SHAPE_TYPE.LINE
 assert p.slides[1].shapes[4].line.width.pt==2
 assert p.slides[1].shapes[-1].has_table
 assert [[c.text for c in row.cells] for row in p.slides[1].shapes[-1].table.rows]==[['施策','担当','確認する成果'],['顧客集中','営業','継続率'],['業務標準化','運用','作業時間']]
 assert p.slides[0].notes_slide.notes_text_frame.text=='この資料は架空の事業データを使った作例です。'
 print('Independent python-pptx: shapes, fonts, lines, table, notes, size passed')
 if args.libreoffice:
  soffice=shutil.which('soffice') or ('/Applications/LibreOffice.app/Contents/MacOS/soffice' if Path('/Applications/LibreOffice.app/Contents/MacOS/soffice').exists() else None)
  if not soffice: raise RuntimeError('LibreOffice is required for this requested check')
  out=temp/'resaved'; out.mkdir()
  command=[soffice,'-env:UserInstallation='+ (temp/'profile').as_uri(),'--headless','--convert-to','pptx:Impress MS PowerPoint 2007 XML','--outdir',str(out)]
  subprocess.run(command+[str(proposal),str(root/'Tests/SwiftSlidesTests/Fixtures/python-pptx.pptx')],check=True,timeout=90)
  for name in ['proposal.pptx','python-pptx.pptx']:
   path=out/name
   assert path.is_file(),'LibreOffice did not generate '+name
   p=Presentation(path); assert len(p.slides)==2
   result=subprocess.run(['swift','run','--package-path',str(root),'swiftslides','text',str(path)],capture_output=True,text=True,check=True)
   assert ('収益性' if name=='proposal.pptx' else '売上成長') in result.stdout
   if name=='proposal.pptx': assert p.slides[1].shapes[-1].has_table
  if args.update_fixture: shutil.copyfile(out/'python-pptx.pptx',root/'Tests/SwiftSlidesTests/Fixtures/libreoffice.pptx')
  print('LibreOffice resave: model and independent text/table checks passed')
