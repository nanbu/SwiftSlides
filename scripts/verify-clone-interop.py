#!/usr/bin/env python3
"""Public consumer + independent clone/import checks; optional mandatory LibreOffice resave."""
import argparse
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile

from lxml import etree
from pptx import Presentation

SWIFT = r'''
import Foundation
import SwiftSlides

let input = URL(filePath: CommandLine.arguments[1])
let output = URL(filePath: CommandLine.arguments[2])
let source = try Presentation(contentsOf: input)
var duplicate = source
try duplicate.duplicateSlide(id: source.slides[0].id, at: 1)
try duplicate.duplicateSlide(id: source.slides[1].id)
let plan = try duplicate.planWrite(options: .init(strict: true))
precondition(plan.canSave && plan.profile == .ooxmlTransitional)
let savedDuplicate = try duplicate.save(to: output.appendingPathComponent("duplicate.pptx"), using: plan)
precondition(savedDuplicate.warnings.isEmpty)
var imported = Presentation(size: source.size)
var local = Slide(name: "取り込み先")
local.addText("取り込み先", frame: .init(x: 20, y: 20, width: 200, height: 50))
imported.slides.append(local)
try imported.importSlide(id: source.slides[0].id, from: source)
try imported.importSlide(id: source.slides[1].id, from: source)
let savedImport = try imported.write(to: output.appendingPathComponent("imported.pptx"))
precondition(savedImport.warnings.isEmpty)
let inventory = try imported.inspectPreservation()
precondition(inventory.parts.isEmpty) // 原本のない新規文書
duplicate.slides[1].notes = .init("複製先の独立したノート")
let savedEdit = try duplicate.write(to: output.appendingPathComponent("edited-clone.pptx"))
print(savedEdit.warnings)
print("公開consumer: inventory・保存計画・複製・取り込みを検証")
'''


def check_relationships(path):
    with zipfile.ZipFile(path) as archive:
        parts = set(archive.namelist())
        master_ids, layout_ids = [], []
        notes_masters = 0
        for part in sorted(parts):
            if part.endswith('.xml'):
                node = etree.fromstring(archive.read(part))
                if etree.QName(node).localname == 'notesMaster':
                    notes_masters += 1
                master_ids += node.xpath('//*[local-name()="sldMasterId"]/@id')
                layout_ids += node.xpath('//*[local-name()="sldLayoutId"]/@id')
            if part.endswith('.rels'):
                node = etree.fromstring(archive.read(part))
                ids = [child.get('Id') for child in node]
                assert len(ids) == len(set(ids)), 'duplicate relationship IDs'
                source_dir = Path(part).parent.parent
                for child in node:
                    if child.get('TargetMode') == 'External':
                        continue
                    target = child.get('Target')
                    if target.startswith('/'):
                        resolved = target[1:]
                    else:
                        # PurePosixPath does not collapse '..'. Avoid filesystem resolution.
                        segments = list(source_dir.parts) if str(source_dir) != '.' else []
                        for segment in target.split('/'):
                            if segment == '..': segments.pop()
                            elif segment not in ('', '.'): segments.append(segment)
                        resolved = '/'.join(segments)
                    assert resolved in parts, 'missing dependency: ' + resolved
        assert notes_masters <= 1
        assert len(master_ids) == len(set(master_ids))
        assert len(layout_ids) == len(set(layout_ids))


def verify(folder):
    duplicate = Presentation(folder / 'duplicate.pptx')
    assert len(duplicate.slides) == 4
    assert duplicate.slides[0].shapes[0].text == duplicate.slides[1].shapes[0].text
    assert duplicate.slides[0].notes_slide.notes_text_frame.text == duplicate.slides[1].notes_slide.notes_text_frame.text
    assert duplicate.slides[0].shapes[5].image.blob == duplicate.slides[1].shapes[5].image.blob
    chart = duplicate.slides[2].shapes[0].chart
    cloned_chart = duplicate.slides[3].shapes[0].chart
    assert chart.part.partname != cloned_chart.part.partname
    assert chart.part.blob == cloned_chart.part.blob
    assert chart.part.chart_workbook.xlsx_part.partname != cloned_chart.part.chart_workbook.xlsx_part.partname
    assert chart.part.chart_workbook.xlsx_part.blob == cloned_chart.part.chart_workbook.xlsx_part.blob
    imported = Presentation(folder / 'imported.pptx')
    assert len(imported.slides) == 3 and len(imported.slide_masters) == 2
    assert imported.slides[0].shapes[0].text == '取り込み先'
    assert imported.slides[1].shapes[0].text == duplicate.slides[0].shapes[0].text
    assert imported.slides[1].shapes[5].image.blob == duplicate.slides[0].shapes[5].image.blob
    assert imported.slides[2].shapes[0].chart.series[0].values == chart.series[0].values
    assert imported.slides[1].slide_layout.slide_master.part.partname != imported.slides[0].slide_layout.slide_master.part.partname
    edited = Presentation(folder / 'edited-clone.pptx')
    assert edited.slides[0].notes_slide.notes_text_frame.text == duplicate.slides[0].notes_slide.notes_text_frame.text
    assert edited.slides[1].notes_slide.notes_text_frame.text == '複製先の独立したノート'
    for name in ('duplicate.pptx', 'imported.pptx', 'edited-clone.pptx'):
        check_relationships(folder / name)
    print('独立python-pptx: notes・chart/workbook・画像・master・ID・依存参照を検証')


def verify_schema(folder, directory, original, allow_preserved_errors):
    schema_names = {
        'http://schemas.openxmlformats.org/presentationml/2006/main': 'pml.xsd',
        'http://schemas.openxmlformats.org/drawingml/2006/main': 'dml-main.xsd',
        'http://schemas.openxmlformats.org/drawingml/2006/chart': 'dml-chart.xsd',
    }
    schemas = {namespace: etree.XMLSchema(etree.parse(str(directory / name))) for namespace, name in schema_names.items()}
    count = 0
    inherited = []
    with zipfile.ZipFile(original) as archive:
        source_parts = {archive.read(part) for part in archive.namelist() if part.endswith('.xml')}
    for name in ('duplicate.pptx', 'imported.pptx', 'edited-clone.pptx'):
        with zipfile.ZipFile(folder / name) as archive:
            for part in archive.namelist():
                if not part.endswith('.xml'): continue
                node = etree.fromstring(archive.read(part))
                namespace = etree.QName(node).namespace
                if namespace in schemas:
                    try:
                        schemas[namespace].assertValid(node)
                        count += 1
                    except etree.DocumentInvalid as error:
                        if not allow_preserved_errors or archive.read(part) not in source_parts:
                            raise
                        inherited.append((name, part, str(error)))
    assert count > 0, '空のschema検査集合'
    print('OOXML XSD: ' + str(count) + 'パーツが適合、原本bytes一致のスキーマ違反は' + str(len(inherited)) + 'パーツ（適合に数えない）')
    for name, part, error in inherited:
        print('原本由来のスキーマ違反: ' + name + ' / ' + part + ': ' + error)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--libreoffice', action='store_true')
    parser.add_argument('--schema-dir', type=Path, help='公開OOXML XSDのローカル配置先（同梱しない）')
    parser.add_argument('--allow-preserved-schema-errors', action='store_true', help='原本と全bytes一致のパーツだけ既存スキーマ違反の保持を許可し別集計')
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    with tempfile.TemporaryDirectory(prefix='swiftslides-clone-oracle-') as directory:
        temp = Path(directory)
        (temp / 'Sources/CloneOracle').mkdir(parents=True)
        (temp / 'Sources/CloneOracle/main.swift').write_text(SWIFT)
        (temp / 'Package.swift').write_text('// swift-tools-version: 6.4\nimport PackageDescription\n'
            'let package = Package(name: "CloneOracle", platforms: [.macOS(.v14), .iOS(.v17)], dependencies: [.package(path: ' + json.dumps(str(root)) + ')], '
            'targets: [.executableTarget(name: "CloneOracle", dependencies: [.product(name: "SwiftSlides", package: "SwiftSlides")])], '
            'swiftLanguageModes: [.v6])\n')
        subprocess.run(['swift', 'run', '--package-path', str(temp), '--scratch-path', str(root / '.build/clone-interop'),
                        'CloneOracle', str(root / 'Tests/SwiftSlidesTests/Fixtures/python-pptx.pptx'), str(temp)], check=True)
        verify(temp)
        if args.schema_dir:
            verify_schema(temp, args.schema_dir, root / 'Tests/SwiftSlidesTests/Fixtures/python-pptx.pptx', args.allow_preserved_schema_errors)
        if args.libreoffice:
            soffice = shutil.which('soffice')
            installed = Path('/Applications/LibreOffice.app/Contents/MacOS/soffice')
            if not soffice and installed.exists(): soffice = str(installed)
            if not soffice: raise RuntimeError('要求されたLibreOffice検査には実アプリが必要です')
            resaved = temp / 'resaved'; resaved.mkdir()
            subprocess.run([soffice, '-env:UserInstallation=' + (temp / 'profile').as_uri(), '--headless', '--convert-to',
                'pptx:Impress MS PowerPoint 2007 XML', '--outdir', str(resaved)]
                + [str(temp / name) for name in ('duplicate.pptx', 'imported.pptx', 'edited-clone.pptx')], check=True, timeout=90)
            # Producer may reorganize parts/masters; compare the semantic properties independently.
            for name, count in [('duplicate.pptx', 4), ('imported.pptx', 3), ('edited-clone.pptx', 4)]:
                p = Presentation(resaved / name)
                assert len(p.slides) == count
                assert any(shape.has_chart for slide in p.slides for shape in slide.shapes)
                assert '売上成長' in p.slides[1].shapes[0].text
                assert p.slides[1].notes_slide.notes_text_frame.text == ('複製先の独立したノート' if name == 'edited-clone.pptx' else Presentation(temp / name).slides[1].notes_slide.notes_text_frame.text)
            print('LibreOffice再保存: 複製・取り込み・独立ノートの3文書を検証')


if __name__ == '__main__':
    main()
